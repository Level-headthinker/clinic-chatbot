"""Bulk data import with a 3-step column-mapping flow.

Clinics' spreadsheets never match our field names, so every import goes:

  Step 1  POST /import/{entity}/preview
          Upload the file → we return its headers, the first 5 data rows, a
          fuzzy-matched suggested mapping (rapidfuzz; ≥80% confidence is
          auto-applied, below that it's a suggestion the user must confirm),
          the clinic's saved mapping from last time, and which fields are
          required. NOTHING is written to the database.

  Step 2  (frontend) The user reviews/edits the mapping — their column on the
          left, a dropdown of our fields (or "Skip") on the right. Required
          fields must be mapped to proceed.

  Step 3  POST /import/{entity}/commit  with the mapping JSON
          dry_run=true  → validate everything, return row-by-row errors only.
          dry_run=false → import the valid rows, return a per-row error report
          plus a CSV of failed rows (base64) so the clinic can fix and
          re-import just those. The mapping is saved per clinic per entity so
          next upload with the same headers is pre-mapped.

Files: CSV or Excel (.xlsx). Max 5 MB / 2000 rows per upload.
"""
from __future__ import annotations

import base64
import csv
import io
import re
import secrets
from datetime import datetime
from typing import Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from sqlalchemy.orm import Session

from app.database import get_db
from app.models.appointment import Appointment
from app.models.doctor import Doctor
from app.models.import_mapping import ImportMapping
from app.models.patient import Patient
from app.models.service import Service
from app.models.user import User
from app.services.auth import hash_password, require_admin_user

try:
    from rapidfuzz import fuzz
    def _similarity(a: str, b: str) -> float:
        return fuzz.token_sort_ratio(a, b)
except ImportError:  # pragma: no cover — rapidfuzz is in requirements
    from difflib import SequenceMatcher
    def _similarity(a: str, b: str) -> float:
        return SequenceMatcher(None, a, b).ratio() * 100

router = APIRouter(prefix="/import", tags=["Data Import"])

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
AUTO_APPLY_CONFIDENCE = 80.0

# ── Field specs per entity ─────────────────────────────────────────────────────
# label/aliases feed the fuzzy matcher; "required" gates step 3.

ENTITY_FIELDS: dict[str, dict[str, dict]] = {
    "patients": {
        "name":              {"label": "Patient Name", "required": True,
                              "aliases": ["patient name", "full name", "name of patient", "patient"]},
        "phone":             {"label": "Phone Number", "required": True, "type": "phone",
                              "aliases": ["phone number", "mobile", "contact", "cell", "whatsapp"]},
        "age":               {"label": "Age", "type": "int", "aliases": ["age", "patient age", "years"]},
        "gender":            {"label": "Gender", "aliases": ["gender", "sex"]},
        "blood_group":       {"label": "Blood Group", "aliases": ["blood group", "blood type"]},
        "allergies":         {"label": "Allergies", "aliases": ["allergies", "allergy"]},
        "chronic_conditions": {"label": "Chronic Conditions",
                               "aliases": ["chronic conditions", "conditions", "medical history", "history"]},
        "emergency_contact": {"label": "Emergency Contact",
                              "aliases": ["emergency contact", "emergency number", "guardian phone"]},
    },
    "doctors": {
        "name":          {"label": "Doctor Name", "required": True,
                          "aliases": ["doctor name", "name", "doctor", "dr name", "physician"]},
        "specialty":     {"label": "Specialty", "required": True,
                          "aliases": ["specialty", "speciality", "specialization", "department", "field"]},
        "qualification": {"label": "Qualification", "aliases": ["qualification", "degree", "degrees"]},
        "fee":           {"label": "Consultation Fee", "aliases": ["fee", "fees", "consultation fee", "charges", "price"]},
        "bio":           {"label": "Bio", "aliases": ["bio", "about", "profile", "description"]},
        "treatments":    {"label": "Treatments (comma-separated)",
                          "aliases": ["treatments", "services", "procedures"]},
    },
    "staff": {
        "full_name": {"label": "Full Name", "required": True,
                      "aliases": ["full name", "name", "staff name", "employee name"]},
        "email":     {"label": "Email", "required": True, "type": "email",
                      "aliases": ["email", "email address", "e-mail"]},
        "phone":     {"label": "Phone", "type": "phone", "aliases": ["phone", "mobile", "contact"]},
    },
    "services": {
        "name":             {"label": "Service Name", "required": True,
                             "aliases": ["service name", "name", "service", "treatment", "procedure"]},
        "price":            {"label": "Price", "type": "number", "aliases": ["price", "fee", "cost", "charges", "amount"]},
        "duration_minutes": {"label": "Duration (minutes)", "type": "int",
                             "aliases": ["duration", "minutes", "time", "duration minutes"]},
    },
    "appointments": {
        "patient_name":  {"label": "Patient Name", "required": True,
                          "aliases": ["patient name", "name", "patient"]},
        "patient_phone": {"label": "Patient Phone", "required": True, "type": "phone",
                          "aliases": ["patient phone", "phone", "mobile", "contact"]},
        "doctor_name":   {"label": "Doctor Name", "required": True,
                          "aliases": ["doctor name", "doctor", "dr", "physician", "consultant"]},
        "slot_datetime": {"label": "Date & Time", "required": True, "type": "datetime",
                          "aliases": ["date and time", "datetime", "appointment date", "date", "slot", "schedule"]},
        "status":        {"label": "Status", "aliases": ["status", "state"]},
        "patient_concern": {"label": "Concern / Reason",
                            "aliases": ["concern", "reason", "complaint", "notes", "purpose"]},
    },
}


def _entity_or_400(entity: str) -> dict:
    spec = ENTITY_FIELDS.get(entity)
    if not spec:
        raise HTTPException(
            status_code=400,
            detail=f"Unknown import type. Use one of: {', '.join(ENTITY_FIELDS)}",
        )
    return spec


# ── File parsing ───────────────────────────────────────────────────────────────

async def _read_table(file: UploadFile) -> tuple[list[str], list[list[str]]]:
    """Parse CSV/XLSX into (headers, rows-of-strings). Validates size and type."""
    raw = await file.read()
    if len(raw) > MAX_FILE_BYTES:
        raise HTTPException(status_code=413, detail="File too large (max 5 MB).")
    if not raw:
        raise HTTPException(status_code=400, detail="The file is empty.")

    name = (file.filename or "").lower()
    if name.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            ws = wb.active
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append(["" if c is None else str(c).strip() for c in row])
                if len(rows) > MAX_ROWS + 1:
                    raise HTTPException(status_code=413, detail=f"Too many rows (max {MAX_ROWS}).")
        except HTTPException:
            raise
        except Exception:
            raise HTTPException(status_code=400, detail="Could not read that Excel file. Is it a valid .xlsx?")
    elif name.endswith(".csv") or name.endswith(".txt"):
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = raw.decode("cp1252")
            except UnicodeDecodeError:
                raise HTTPException(status_code=400, detail="Could not decode the CSV. Save it as UTF-8 and retry.")
        reader = csv.reader(io.StringIO(text))
        rows = [[c.strip() for c in row] for row in reader]
        if len(rows) > MAX_ROWS + 1:
            raise HTTPException(status_code=413, detail=f"Too many rows (max {MAX_ROWS}).")
    else:
        raise HTTPException(status_code=400, detail="Unsupported file type. Upload a .csv or .xlsx file.")

    rows = [r for r in rows if any(c for c in r)]
    if len(rows) < 2:
        raise HTTPException(status_code=400, detail="The file needs a header row plus at least one data row.")
    headers = [h or f"Column {i + 1}" for i, h in enumerate(rows[0])]
    return headers, rows[1:]


# ── Fuzzy mapping suggestion ───────────────────────────────────────────────────

def _normalize_header(h: str) -> str:
    # "PatientName" → "patient name"; strip punctuation/underscores
    h = re.sub(r"(?<=[a-z0-9])(?=[A-Z])", " ", h)
    return re.sub(r"[^a-z0-9 ]+", " ", h.lower()).strip()


def suggest_mapping(headers: list[str], spec: dict) -> dict:
    """For each file column, the best-matching target field with confidence.
    ≥80 → auto-applied; below → suggestion only (frontend shows it unselected)."""
    out = {}
    for header in headers:
        norm = _normalize_header(header)
        squashed = norm.replace(" ", "")
        best_field, best_score = None, 0.0
        for field, meta in spec.items():
            candidates = [field.replace("_", " "), meta["label"].lower()] + meta.get("aliases", [])
            # Compare both token-wise and with spaces squashed, so
            # "WhatsApp" ↔ "whatsapp" and "PatientName" ↔ "patient name" both score high.
            score = max(
                max(_similarity(norm, c.lower()),
                    _similarity(squashed, c.lower().replace(" ", "")))
                for c in candidates
            )
            if score > best_score:
                best_field, best_score = field, score
        out[header] = {
            "field": best_field if best_score >= 50 else None,
            "confidence": round(best_score, 1),
            "auto": best_score >= AUTO_APPLY_CONFIDENCE,
        }
    return out


# ── Value validation / coercion ────────────────────────────────────────────────

_PK_PHONE = re.compile(r"^(\+?92|0)?3[0-9]{9}$")


def _clean_phone(value: str) -> Optional[str]:
    digits = re.sub(r"[ \-()]", "", value)
    if not _PK_PHONE.match(digits):
        return None
    digits = digits.lstrip("+")
    if digits.startswith("92"):
        digits = "0" + digits[2:]
    elif not digits.startswith("0"):
        digits = "0" + digits
    return digits


_DATE_FORMATS = (
    "%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S",
    "%d/%m/%Y %H:%M", "%d-%m-%Y %H:%M", "%d/%m/%Y %I:%M %p", "%d-%m-%Y %I:%M %p",
    "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y",
)


def _parse_datetime(value: str) -> Optional[datetime]:
    value = value.strip()
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(value)
    except ValueError:
        return None


_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def _coerce(value: str, ftype: str, label: str) -> tuple[object, Optional[str]]:
    """(coerced_value, plain-language error or None)."""
    value = (value or "").strip()
    if not value:
        return None, None
    if ftype == "phone":
        phone = _clean_phone(value)
        return phone, None if phone else f"'{value}' is not a valid Pakistani mobile number (e.g. 03001234567)"
    if ftype == "int":
        try:
            return int(float(value)), None
        except ValueError:
            return None, f"'{value}' should be a whole number for {label}"
    if ftype == "number":
        try:
            return float(value), None
        except ValueError:
            return None, f"'{value}' should be a number for {label}"
    if ftype == "email":
        return (value.lower(), None) if _EMAIL_RE.match(value) else \
               (None, f"'{value}' is not a valid email address")
    if ftype == "datetime":
        dt = _parse_datetime(value)
        return dt, None if dt else \
            f"'{value}' is not a recognizable date/time (try 2026-06-15 10:30 or 15/06/2026 10:30)"
    return value, None


def _apply_mapping_and_validate(
    headers: list[str], rows: list[list[str]], mapping: dict, spec: dict
) -> tuple[list[dict], list[dict]]:
    """→ (valid_records, errors). Each error: {row: N, errors: [msgs]} with N
    being the spreadsheet row number (header = row 1)."""
    field_by_index: dict[int, str] = {}
    for idx, header in enumerate(headers):
        field = mapping.get(header)
        if field and field in spec:
            field_by_index[idx] = field

    mapped_fields = set(field_by_index.values())
    missing = [
        spec[f]["label"] for f, meta in spec.items()
        if meta.get("required") and f not in mapped_fields
    ]
    if missing:
        raise HTTPException(
            status_code=400,
            detail=f"Required field(s) not mapped: {', '.join(missing)}. "
                   "Map them in the column-mapping step before importing.",
        )

    valid, errors = [], []
    for i, row in enumerate(rows):
        record: dict = {}
        row_errors: list[str] = []
        errored_fields: set[str] = set()
        for idx, field in field_by_index.items():
            meta = spec[field]
            raw = row[idx] if idx < len(row) else ""
            value, err = _coerce(raw, meta.get("type", "str"), meta["label"])
            if err:
                row_errors.append(err)
                errored_fields.add(field)
            elif value is not None:
                record[field] = value
        for field, meta in spec.items():
            if meta.get("required") and record.get(field) in (None, "") \
                    and field not in errored_fields:
                row_errors.append(f"{meta['label']} is missing")
        if row_errors:
            errors.append({"row": i + 2, "errors": row_errors})
        else:
            valid.append(record)
    return valid, errors


def _failed_rows_csv(headers: list[str], rows: list[list[str]], errors: list[dict]) -> str:
    """CSV (base64) of only the failed rows + an extra column explaining why."""
    bad_indexes = {e["row"] - 2 for e in errors}
    reasons = {e["row"] - 2: "; ".join(e["errors"]) for e in errors}
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers + ["IMPORT ERRORS — fix these and re-upload"])
    for i in sorted(bad_indexes):
        if 0 <= i < len(rows):
            writer.writerow(list(rows[i]) + [reasons.get(i, "")])
    return base64.b64encode(buf.getvalue().encode("utf-8-sig")).decode()


# ── Row → DB record creators (per entity) ──────────────────────────────────────

def _import_rows(entity: str, records: list[dict], db: Session, user: User) -> tuple[int, int, list[dict]]:
    """Insert records for the user's tenant. → (created, skipped_duplicates, extra_errors)."""
    tenant_id = user.tenant_id
    branch_id = user.branch_id  # branch admins import into their own branch
    created = skipped = 0
    extra_errors: list[dict] = []

    if entity == "patients":
        existing_phones = {
            p[0] for p in db.query(Patient.phone).filter(Patient.tenant_id == tenant_id).all()
        }
        for r in records:
            if r["phone"] in existing_phones:
                skipped += 1
                continue
            db.add(Patient(tenant_id=tenant_id, primary_branch_id=branch_id, **r))
            existing_phones.add(r["phone"])
            created += 1

    elif entity == "doctors":
        existing = {
            (d[0] or "").strip().lower()
            for d in db.query(Doctor.name).filter(Doctor.tenant_id == tenant_id).all()
        }
        for r in records:
            if r["name"].strip().lower() in existing:
                skipped += 1
                continue
            treatments = r.pop("treatments", None)
            if isinstance(treatments, str):
                treatments = [t.strip() for t in treatments.split(",") if t.strip()]
            fee = r.pop("fee", None)
            db.add(Doctor(
                tenant_id=tenant_id, branch_id=branch_id,
                treatments=treatments or [], fee=str(fee) if fee is not None else None,
                **r,
            ))
            existing.add(r["name"].strip().lower())
            created += 1

    elif entity == "staff":
        existing = {
            (u[0] or "").lower()
            for u in db.query(User.email).all()  # emails are globally unique
        }
        for r in records:
            if r["email"] in existing:
                skipped += 1
                continue
            temp_password = secrets.token_urlsafe(9)
            db.add(User(
                tenant_id=tenant_id, branch_id=branch_id,
                email=r["email"], full_name=r["full_name"],
                role="staff", hashed_password=hash_password(temp_password),
            ))
            existing.add(r["email"])
            created += 1
            # Surface the temp password to the importing admin (one-time view).
            extra_errors.append({
                "row": None,
                "info": f"{r['email']} created — temporary password: {temp_password}",
            })

    elif entity == "services":
        existing = {
            (s[0] or "").strip().lower()
            for s in db.query(Service.name).filter(Service.tenant_id == tenant_id).all()
        }
        for r in records:
            if r["name"].strip().lower() in existing:
                skipped += 1
                continue
            db.add(Service(
                tenant_id=tenant_id, branch_id=branch_id,
                name=r["name"], price=r.get("price"),
                duration_minutes=r.get("duration_minutes") or 30,
            ))
            existing.add(r["name"].strip().lower())
            created += 1

    elif entity == "appointments":
        doctors = db.query(Doctor).filter(Doctor.tenant_id == tenant_id).all()
        by_name = {(d.name or "").strip().lower(): d for d in doctors}
        for idx, r in enumerate(records):
            wanted = r["doctor_name"].strip().lower().removeprefix("dr.").removeprefix("dr").strip()
            doctor = by_name.get(r["doctor_name"].strip().lower()) or next(
                (d for n, d in by_name.items() if wanted and wanted in n), None
            )
            if not doctor:
                extra_errors.append({
                    "row": None,
                    "errors": [f"Doctor '{r['doctor_name']}' not found in your clinic — add the doctor first"],
                })
                skipped += 1
                continue
            status_val = (r.get("status") or "completed").strip().lower().replace("-", "_")
            dup = db.query(Appointment).filter(
                Appointment.tenant_id == tenant_id,
                Appointment.doctor_id == doctor.id,
                Appointment.slot_datetime == r["slot_datetime"],
                Appointment.patient_phone == r["patient_phone"],
            ).first()
            if dup:
                skipped += 1
                continue
            db.add(Appointment(
                tenant_id=tenant_id, branch_id=branch_id, doctor_id=doctor.id,
                patient_name=r["patient_name"], patient_phone=r["patient_phone"],
                patient_concern=r.get("patient_concern") or "Imported",
                slot_datetime=r["slot_datetime"], status=status_val,
            ))
            created += 1

    return created, skipped, extra_errors


def _save_mapping(db: Session, tenant_id, entity: str, mapping: dict) -> None:
    row = db.query(ImportMapping).filter(
        ImportMapping.tenant_id == tenant_id, ImportMapping.entity == entity
    ).first()
    if row:
        row.mapping = mapping
    else:
        db.add(ImportMapping(tenant_id=tenant_id, entity=entity, mapping=mapping))


# ── Endpoints ──────────────────────────────────────────────────────────────────

@router.get("/{entity}/fields")
def list_fields(entity: str, current_user: User = Depends(require_admin_user)):
    """Target fields for the mapping dropdowns (label, required, type)."""
    spec = _entity_or_400(entity)
    return [
        {"field": f, "label": m["label"], "required": bool(m.get("required")),
         "type": m.get("type", "str")}
        for f, m in spec.items()
    ]


@router.post("/{entity}/preview")
async def preview_import(
    entity: str,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin_user),
):
    """Step 1: read the file, return headers + sample rows + suggested mapping.
    Writes nothing to the database."""
    spec = _entity_or_400(entity)
    headers, rows = await _read_table(file)

    saved = db.query(ImportMapping).filter(
        ImportMapping.tenant_id == current_user.tenant_id,
        ImportMapping.entity == entity,
    ).first()
    saved_mapping = saved.mapping if saved else {}
    suggestions = suggest_mapping(headers, spec)
    # A previously saved mapping for the same header wins over fuzzy guessing.
    prefill = {
        h: (saved_mapping.get(h) if saved_mapping.get(h) in spec
            else (s["field"] if s["auto"] else None))
        for h, s in suggestions.items()
    }

    return {
        "entity": entity,
        "headers": headers,
        "sample_rows": rows[:5],
        "total_data_rows": len(rows),
        "suggestions": suggestions,   # full detail incl. confidence for the UI
        "prefill_mapping": prefill,   # what the mapping screen starts with
        "required_fields": [f for f, m in spec.items() if m.get("required")],
        "fields": [
            {"field": f, "label": m["label"], "required": bool(m.get("required"))}
            for f, m in spec.items()
        ],
    }


@router.post("/{entity}/commit")
async def commit_import(
    entity: str,
    file: UploadFile = File(...),
    mapping: str = Form(...),          # JSON: {"Their Column": "our_field" | null}
    dry_run: bool = Form(False),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin_user),
):
    """Step 3: apply the mapping, validate every row, then import valid rows
    (dry_run=true validates only). Returns a row-by-row error report and a
    base64 CSV of failed rows for fix-and-reimport."""
    import json
    spec = _entity_or_400(entity)
    try:
        mapping_dict = json.loads(mapping)
        if not isinstance(mapping_dict, dict):
            raise ValueError
    except ValueError:
        raise HTTPException(status_code=400, detail="mapping must be a JSON object")
    mapping_dict = {k: v for k, v in mapping_dict.items() if v}

    headers, rows = await _read_table(file)
    valid, errors = _apply_mapping_and_validate(headers, rows, mapping_dict, spec)

    result = {
        "entity": entity,
        "total_rows": len(rows),
        "valid_rows": len(valid),
        "failed_rows": len(errors),
        "errors": errors[:200],
        "dry_run": dry_run,
        "created": 0,
        "skipped_duplicates": 0,
    }
    if errors:
        result["failed_rows_csv_base64"] = _failed_rows_csv(headers, rows, errors)

    if dry_run:
        return result

    try:
        created, skipped, extra = _import_rows(entity, valid, db, current_user)
        _save_mapping(db, current_user.tenant_id, entity, mapping_dict)
        db.commit()
    except HTTPException:
        db.rollback()
        raise
    except Exception:
        db.rollback()
        raise HTTPException(
            status_code=500,
            detail="Import failed — nothing was saved. Please try again.",
        )
    result["created"] = created
    result["skipped_duplicates"] = skipped
    if extra:
        result["notes"] = extra
    return result
