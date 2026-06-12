"""Clinic & platform performance reports.

Builds daily / weekly / monthly / yearly metric snapshots per clinic (or
platform-wide when ``tenant_id`` is None), stores them in ``system_reports``
(clinic_id column distinguishes the two), and renders branded XLSX / PDF
exports that are cached on disk for re-download.

Metrics per period:
  - messages: chat sessions started + user messages stored for those sessions,
    plus WhatsApp usage counters (message timestamps are not stored per-message,
    so message counts are attributed to the session's creation period)
  - appointments booked / completed / no-show / cancelled
  - busiest doctor, busiest time slot (hour of day)
  - new vs returning patients (first-ever appointment inside vs before period)
  - revenue summary from invoices (billed + collected)
  - leads captured
"""
from __future__ import annotations

import logging
import os
from datetime import datetime, timedelta, timezone

from sqlalchemy import func

from app.models.appointment import Appointment
from app.models.chat import ChatSession, Lead
from app.models.doctor import Doctor
from app.models.invoice import Invoice
from app.models.system_report import SystemReport
from app.models.tenant import Tenant
from app.models.whatsapp_number import WhatsAppNumberMapping

logger = logging.getLogger(__name__)

REPORT_TYPES = ("daily", "weekly", "monthly", "yearly")

# Export files live outside the repo tree's tracked files; created on demand.
REPORTS_DIR = os.environ.get(
    "REPORTS_DIR",
    os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), "reports_store"),
)


def period_bounds(report_type: str, ref: datetime | None = None) -> tuple[datetime, datetime]:
    """[start, end) bounds of the period CONTAINING ref (default: now, UTC)."""
    now = ref or datetime.now(timezone.utc)
    if now.tzinfo is None:
        now = now.replace(tzinfo=timezone.utc)
    day = now.replace(hour=0, minute=0, second=0, microsecond=0)
    if report_type == "daily":
        return day, day + timedelta(days=1)
    if report_type == "weekly":
        start = day - timedelta(days=day.weekday())  # Monday
        return start, start + timedelta(days=7)
    if report_type == "monthly":
        start = day.replace(day=1)
        nxt = (start + timedelta(days=32)).replace(day=1)
        return start, nxt
    if report_type == "yearly":
        start = day.replace(month=1, day=1)
        return start, start.replace(year=start.year + 1)
    raise ValueError(f"report_type must be one of {REPORT_TYPES}")


def previous_period_bounds(report_type: str, ref: datetime | None = None):
    """Bounds of the COMPLETED period before ref — what the cron reports on."""
    start, _ = period_bounds(report_type, ref)
    return period_bounds(report_type, start - timedelta(seconds=1))


def _hour_label(hour: int) -> str:
    dt = datetime(2000, 1, 1, hour)
    return dt.strftime("%I %p").lstrip("0")


def build_clinic_metrics(db, tenant_id, start: datetime, end: datetime) -> dict:
    """All metrics for one clinic in [start, end). tenant_id=None → platform-wide."""
    def scoped(q, model):
        return q.filter(model.tenant_id == tenant_id) if tenant_id else q

    # ── Messages / sessions ──────────────────────────────────
    sessions = scoped(db.query(ChatSession), ChatSession).filter(
        ChatSession.created_at >= start, ChatSession.created_at < end
    ).all()
    user_messages = sum(
        1 for s in sessions for m in (s.messages or []) if m.get("role") == "user"
    )

    # ── Appointments ─────────────────────────────────────────
    appts_q = scoped(db.query(Appointment), Appointment).filter(
        Appointment.created_at >= start, Appointment.created_at < end
    )
    appts = appts_q.all()
    by_status: dict[str, int] = {}
    for a in appts:
        key = (a.status or "unknown").lower().replace("-", "_")
        by_status[key] = by_status.get(key, 0) + 1

    # Busiest doctor
    doctor_counts: dict = {}
    for a in appts:
        doctor_counts[a.doctor_id] = doctor_counts.get(a.doctor_id, 0) + 1
    busiest_doctor = None
    if doctor_counts:
        top_id = max(doctor_counts, key=doctor_counts.get)
        doc = db.query(Doctor).filter(Doctor.id == top_id).first()
        if doc:
            name = doc.name if (doc.name or "").lower().startswith("dr") else f"Dr. {doc.name}"
            busiest_doctor = {"name": name, "appointments": doctor_counts[top_id]}

    # Busiest slot (hour of day)
    hour_counts: dict[int, int] = {}
    for a in appts:
        if a.slot_datetime:
            hour_counts[a.slot_datetime.hour] = hour_counts.get(a.slot_datetime.hour, 0) + 1
    busiest_slot = None
    if hour_counts:
        top_hour = max(hour_counts, key=hour_counts.get)
        busiest_slot = {
            "slot": f"{_hour_label(top_hour)} – {_hour_label((top_hour + 1) % 24)}",
            "appointments": hour_counts[top_hour],
        }

    # New vs returning patients (by phone, first appointment ever)
    phones = {a.patient_phone for a in appts if a.patient_phone}
    new_patients = returning_patients = 0
    for phone in phones:
        before = scoped(db.query(Appointment), Appointment).filter(
            Appointment.patient_phone == phone,
            Appointment.created_at < start,
        ).first()
        if before:
            returning_patients += 1
        else:
            new_patients += 1

    # Daily appointment breakdown (for the usage graph / report table)
    daily: dict[str, int] = {}
    for a in appts:
        if a.created_at:
            daily_key = a.created_at.strftime("%Y-%m-%d")
            daily[daily_key] = daily.get(daily_key, 0) + 1

    # ── Revenue ──────────────────────────────────────────────
    inv_q = scoped(db.query(
        func.coalesce(func.sum(Invoice.total_amount), 0),
        func.coalesce(func.sum(Invoice.paid_amount), 0),
        func.count(Invoice.id),
    ), Invoice).filter(Invoice.created_at >= start, Invoice.created_at < end)
    billed, collected, invoice_count = inv_q.one()

    # ── Leads ────────────────────────────────────────────────
    leads = scoped(db.query(func.count(Lead.id)), Lead).filter(
        Lead.created_at >= start, Lead.created_at < end
    ).scalar() or 0

    # ── WhatsApp usage (current counters) ────────────────────
    wa_q = db.query(
        func.coalesce(func.sum(WhatsAppNumberMapping.messages_used_this_month), 0),
        func.coalesce(func.sum(WhatsAppNumberMapping.message_limit_monthly), 0),
    )
    if tenant_id:
        wa_q = wa_q.filter(WhatsAppNumberMapping.tenant_id == tenant_id)
    wa_used, wa_limit = wa_q.one()

    return {
        "chat_sessions_started": len(sessions),
        "messages_received": user_messages,
        "whatsapp_used_this_month": int(wa_used),
        "whatsapp_limit_monthly": int(wa_limit),
        "appointments_booked": len(appts),
        "appointments_by_status": by_status,
        "appointments_completed": by_status.get("completed", 0),
        "appointments_no_show": by_status.get("no_show", 0),
        "appointments_cancelled": by_status.get("cancelled", 0),
        "busiest_doctor": busiest_doctor,
        "busiest_slot": busiest_slot,
        "new_patients": new_patients,
        "returning_patients": returning_patients,
        "leads_captured": int(leads),
        "revenue": {
            "invoices": int(invoice_count),
            "billed": int(billed),
            "collected": int(collected),
        },
        "daily_appointments": dict(sorted(daily.items())),
    }


def generate_report(
    db,
    tenant_id,
    report_type: str,
    start: datetime,
    end: datetime,
    save: bool = True,
) -> SystemReport:
    metrics = build_clinic_metrics(db, tenant_id, start, end)
    report = SystemReport(
        clinic_id=tenant_id,
        report_type=report_type,
        period_start=start,
        period_end=end,
        metrics=metrics,
    )
    if save:
        db.add(report)
        db.commit()
        db.refresh(report)
    return report


# ════════════════════════════════════════════════════════════
# EXPORTS — branded XLSX and PDF, cached on disk
# ════════════════════════════════════════════════════════════

_ROWS = [
    ("Chat sessions started", "chat_sessions_started"),
    ("Messages received", "messages_received"),
    ("WhatsApp used (this month)", "whatsapp_used_this_month"),
    ("Appointments booked", "appointments_booked"),
    ("Appointments completed", "appointments_completed"),
    ("No-shows", "appointments_no_show"),
    ("Cancelled", "appointments_cancelled"),
    ("New patients", "new_patients"),
    ("Returning patients", "returning_patients"),
    ("Leads captured", "leads_captured"),
]


def _flat_rows(metrics: dict) -> list[tuple[str, str]]:
    rows = [(label, str(metrics.get(key, 0))) for label, key in _ROWS]
    bd = metrics.get("busiest_doctor")
    rows.append(("Busiest doctor",
                 f"{bd['name']} ({bd['appointments']} appts)" if bd else "—"))
    bs = metrics.get("busiest_slot")
    rows.append(("Busiest time slot",
                 f"{bs['slot']} ({bs['appointments']} appts)" if bs else "—"))
    rev = metrics.get("revenue") or {}
    rows.append(("Revenue billed (PKR)", str(rev.get("billed", 0))))
    rows.append(("Revenue collected (PKR)", str(rev.get("collected", 0))))
    rows.append(("Invoices issued", str(rev.get("invoices", 0))))
    return rows


def _export_basename(report: SystemReport, clinic_name: str) -> str:
    safe = "".join(c if c.isalnum() or c in "-_" else "_" for c in clinic_name)[:40]
    return (
        f"{safe}_{report.report_type}_"
        f"{report.period_start.strftime('%Y%m%d')}_{str(report.id)[:8]}"
    )


def export_xlsx(report: SystemReport, clinic_name: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    os.makedirs(REPORTS_DIR, exist_ok=True)
    path = os.path.join(REPORTS_DIR, _export_basename(report, clinic_name) + ".xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = f"{report.report_type.title()} Report"
    ws["A1"] = clinic_name
    ws["A1"].font = Font(size=16, bold=True)
    ws["A2"] = (
        f"{report.report_type.title()} report · "
        f"{report.period_start.strftime('%d %b %Y')} – "
        f"{(report.period_end - timedelta(days=1)).strftime('%d %b %Y')}"
    )
    ws.append([])
    ws.append(["Metric", "Value"])
    ws["A4"].font = ws["B4"].font = Font(bold=True)
    for label, value in _flat_rows(report.metrics or {}):
        ws.append([label, value])

    daily = (report.metrics or {}).get("daily_appointments") or {}
    if daily:
        ws.append([])
        ws.append(["Date", "Appointments booked"])
        for d, n in daily.items():
            ws.append([d, n])
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 28
    wb.save(path)
    return path


def export_pdf(report: SystemReport, clinic_name: str, logo_path: str | None = None) -> str:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    os.makedirs(REPORTS_DIR, exist_ok=True)
    path = os.path.join(REPORTS_DIR, _export_basename(report, clinic_name) + ".pdf")

    styles = getSampleStyleSheet()
    story = []
    if logo_path and os.path.exists(logo_path):
        try:
            story.append(Image(logo_path, width=3 * cm, height=3 * cm))
        except Exception:
            logger.warning("Could not embed clinic logo %s", logo_path)
    story.append(Paragraph(clinic_name, styles["Title"]))
    story.append(Paragraph(
        f"{report.report_type.title()} report · "
        f"{report.period_start.strftime('%d %b %Y')} – "
        f"{(report.period_end - timedelta(days=1)).strftime('%d %b %Y')}",
        styles["Normal"],
    ))
    story.append(Spacer(1, 0.6 * cm))

    data = [["Metric", "Value"]] + [list(r) for r in _flat_rows(report.metrics or {})]
    table = Table(data, colWidths=[9 * cm, 7 * cm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e3a5f")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(table)
    story.append(Spacer(1, 0.6 * cm))
    story.append(Paragraph(
        f"Generated {datetime.now(timezone.utc).strftime('%d %b %Y %H:%M UTC')} · ClinicBot",
        styles["Normal"],
    ))
    SimpleDocTemplate(str(path), pagesize=A4).build(story)
    return path


def ensure_exports(db, report: SystemReport) -> SystemReport:
    """Generate (or reuse) the cached XLSX/PDF files for a stored report."""
    clinic_name = "All Clinics (Platform)"
    if report.clinic_id:
        tenant = db.query(Tenant).filter(Tenant.id == report.clinic_id).first()
        clinic_name = tenant.name if tenant else "Clinic"
    changed = False
    if not report.xlsx_path or not os.path.exists(report.xlsx_path):
        report.xlsx_path = export_xlsx(report, clinic_name)
        changed = True
    if not report.pdf_path or not os.path.exists(report.pdf_path):
        report.pdf_path = export_pdf(report, clinic_name)
        changed = True
    if changed:
        db.commit()
    return report


def generate_scheduled_reports(report_type: str) -> None:
    """Cron entrypoint — one report per active clinic plus a platform rollup,
    for the just-completed period, with export files pre-rendered."""
    from app.database import SessionLocal

    db = SessionLocal()
    try:
        start, end = previous_period_bounds(report_type)
        tenants = db.query(Tenant).filter(Tenant.is_active.is_(True)).all()
        for tenant in tenants:
            try:
                existing = db.query(SystemReport).filter(
                    SystemReport.clinic_id == tenant.id,
                    SystemReport.report_type == report_type,
                    SystemReport.period_start == start,
                ).first()
                report = existing or generate_report(db, tenant.id, report_type, start, end)
                ensure_exports(db, report)
            except Exception:
                db.rollback()
                logger.exception("scheduled %s report failed for clinic %s", report_type, tenant.id)
        try:
            existing = db.query(SystemReport).filter(
                SystemReport.clinic_id.is_(None),
                SystemReport.report_type == report_type,
                SystemReport.period_start == start,
            ).first()
            report = existing or generate_report(db, None, report_type, start, end)
            ensure_exports(db, report)
        except Exception:
            db.rollback()
            logger.exception("scheduled platform %s report failed", report_type)
        logger.info("Scheduled %s reports generated for %d clinics", report_type, len(tenants))
    finally:
        db.close()
